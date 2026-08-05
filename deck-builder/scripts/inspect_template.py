#!/usr/bin/env python3
"""Report the structure, palette and fonts of any .pptx / .potx.

    python3 inspect_template.py /path/to/template.pptx

Run this on a user-supplied template (style option 3) before building anything, and
report what it found. Also useful to re-verify the company template if it changes.

The key thing it answers: which layouts exist, what placeholders each has, and which
layouts have no body placeholder (those need a manual textbox).

Note it reads colours and fonts from the *slides, layouts and master shapes*, not from
the theme. Many corporate templates leave the stock Office theme in place and carry the
real brand only in shape formatting — the company template does exactly that.
"""

import argparse
import collections
import os
import re
import sys
import zipfile

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("python-pptx is not installed. Run: pip3 install python-pptx")

EMU_IN = 914400


def inches(v):
    return None if v is None else round(v / EMU_IN, 2)


def scan_xml(path):
    """Count srgbClr values and typefaces across the parts that carry real formatting."""
    colors = collections.Counter()
    fonts = collections.Counter()
    sizes = collections.Counter()
    theme_fonts = []
    theme_accents = []

    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not name.endswith(".xml"):
                continue
            in_content = name.startswith(("ppt/slides/", "ppt/slideLayouts/",
                                          "ppt/slideMasters/"))
            data = z.read(name).decode("utf-8", "replace")
            if in_content:
                colors.update(m.lower() for m in
                              re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', data))
                fonts.update(re.findall(r'typeface="([^"+][^"]*)"', data))
                sizes.update(int(s) // 100 for s in re.findall(r'\ssz="(\d+)"', data))
            elif name.startswith("ppt/theme/"):
                theme_fonts += re.findall(
                    r'<a:(?:majorFont|minorFont)>\s*<a:latin typeface="([^"]+)"', data)
                theme_accents += re.findall(
                    r'<a:accent1>\s*<a:srgbClr val="([0-9A-Fa-f]{6})"', data)
    return colors, fonts, sizes, theme_fonts, theme_accents


GENERIC_WORDS = ("simple", "basic", "plain", "default", "blank", "content",
                 "general", "standard", "custom")


def is_dark_layout(prs, lay, xml):
    """True if the layout's background is dark, so its text has to be light."""
    def dark(hexv):
        r, g, b = (int(hexv[i:i+2], 16) for i in (0, 2, 4))
        return (0.299 * r + 0.587 * g + 0.114 * b) < 128

    bg = re.search(r"<p:bg>.*?</p:bg>", xml, re.S)
    if bg and any(dark(h) for h in
                  re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', bg.group(0))):
        return True
    full = 0.95 * prs.slide_width * prs.slide_height
    for sh in lay.shapes:
        if not all((sh.width, sh.height)) or sh.width * sh.height < full:
            continue
        cols = re.findall(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', sh._element.xml)
        if cols and all(dark(h) for h in cols[:2]):
            return True
    return False


def duplicate_layouts(path, prs):
    """Group layouts by placeholder geometry plus background tone.

    Background *art* is not part of the key: a different photo behind identical title
    and subtitle boxes is the same layout wearing a different picture.

    Background *tone* is, because a dark layout needs light text. 'Content Heading' and
    'Content Heading Dark' have identical geometry and opposite text colours, so treating
    them as one would make the dark variant unreachable.

    Canonical pick = a name that doesn't assert a subject, then the least decorative
    burnt-in wording, then the template's own order. Mirrors layout_groups() in
    build_deck.py, which does the actual redirecting.
    """
    z = zipfile.ZipFile(path)
    groups = collections.defaultdict(list)
    dark = []

    for i, lay in enumerate(prs.slide_layouts):
        phs = tuple(sorted(
            (p.placeholder_format.idx, str(p.placeholder_format.type),
             p.left, p.top, p.width, p.height) for p in lay.placeholders))

        # the decorative words burnt into the layout — often the only visible difference
        xml = z.read(str(lay.part.partname)[1:]).decode("utf-8", "replace")
        words = [t for t in re.findall(r"<a:t>([^<]*)</a:t>", xml)
                 if t.strip() and "Click to edit" not in t
                 and "Click icon" not in t and t.strip() != "‹#›"]

        on_dark = is_dark_layout(prs, lay, xml)
        if on_dark:
            dark.append((i, lay.name))
        groups[(phs, on_dark)].append((i, lay.name, " ".join(words)))

    def rank(t):
        low = t[1].strip().lower()
        return (0 if any(w in low for w in GENERIC_WORDS) else 1, len(t[2]), t[0])

    return [sorted(v, key=rank) for v in groups.values()], len(groups), dark


def zip_health(path):
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    dupes = [n for n, c in collections.Counter(names).items() if c > 1]
    media = [n for n in names if n.startswith("ppt/media/")]
    return dupes, media


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path")
    ap.add_argument("--json", action="store_true", help="emit machine-readable JSON")
    args = ap.parse_args()

    if not os.path.isfile(args.path):
        sys.exit(f"not found: {args.path}")

    prs = Presentation(args.path)
    layouts = list(prs.slide_layouts)
    colors, fonts, sizes, theme_fonts, theme_accents = scan_xml(args.path)
    dupes, media = zip_health(args.path)

    W, H = prs.slide_width, prs.slide_height
    ratio = "16:9" if abs(W / H - 16 / 9) < 0.02 else \
            "4:3" if abs(W / H - 4 / 3) < 0.02 else f"{W/H:.2f}:1"

    if args.json:
        import json
        print(json.dumps({
            "path": args.path,
            "size_in": [inches(W), inches(H)],
            "ratio": ratio,
            "slides": len(prs.slides),
            "layouts": [
                {"index": i, "name": l.name,
                 "placeholders": [
                     {"idx": p.placeholder_format.idx,
                      "type": str(p.placeholder_format.type),
                      "box_in": [inches(p.left), inches(p.top),
                                 inches(p.width), inches(p.height)]}
                     for p in l.placeholders]}
                for i, l in enumerate(layouts)],
            "colors": colors.most_common(20),
            "fonts": fonts.most_common(12),
            "sizes_pt": sorted(sizes, reverse=True),
            "theme_fonts": sorted(set(theme_fonts)),
            "theme_accent1": sorted(set(theme_accents)),
            "media_count": len(media),
            "duplicate_zip_entries": dupes,
        }, ensure_ascii=False, indent=2))
        return

    print(f"\n{os.path.basename(args.path)}")
    print(f"  {os.path.getsize(args.path)/1e6:.1f} MB · "
          f"{inches(W)} × {inches(H)} in ({ratio}) · "
          f"{len(prs.slides)} slides · {len(layouts)} layouts · "
          f"{len(media)} media files")
    if dupes:
        print(f"  !! {len(dupes)} duplicate zip entries — this package is damaged")

    print("\n--- Layouts " + "-" * 50)
    no_body = []
    for i, lay in enumerate(layouts):
        phs = []
        has_body = False
        for p in lay.placeholders:
            t = str(p.placeholder_format.type)
            phs.append(f"{p.placeholder_format.idx}:{t}")
            if "BODY" in t or "OBJECT" in t:
                has_body = True
        print(f"  [{i:2d}] {lay.name:<32} {', '.join(phs) or '(none)'}")
        if not has_body:
            no_body.append(f"[{i}] {lay.name}")

    if no_body:
        print(f"\n  {len(no_body)} layout(s) have NO body/object placeholder — bulleted")
        print("  content on these needs a manually added textbox:")
        for n in no_body:
            print(f"    {n}")

    groups, n_distinct, dark = duplicate_layouts(args.path, prs)
    if dark:
        print("\n--- Dark layouts " + "-" * 45)
        print("  Text on these must be light — build_deck.py flips anything it adds")
        print("  itself to white, and the verifier measures contrast against the")
        print("  real background rather than assuming white:")
        for i, name in dark:
            print(f"    [{i:2d}] {name}")

    dupes = [g for g in groups if len(g) > 1]
    print("\n--- Interchangeable layouts " + "-" * 34)
    print(f"  {len(layouts)} layouts -> {n_distinct} distinct text arrangements. Layouts")
    print("  that put text in the same place are one layout; the background art behind")
    print("  it doesn't count — but a dark background is a real difference, not art.")
    print("  build_deck.py keeps the first of each group:")
    if not dupes:
        print("\n    none — every layout has its own arrangement")
    for group in dupes:
        print()
        for j, (i, name, words) in enumerate(group):
            mark = "KEEP " if j == 0 else "  dup"
            words = words if len(words) <= 44 else words[:41] + "..."
            print(f"    {mark} [{i:2d}] {name:<30} art says: {words or '—'!r}")

    special = []
    for i, lay in enumerate(layouts):
        low = lay.name.strip().lower()
        if low in ("thank you", "thankyou", "thanks", "closing"):
            extra = [sh.text_frame.text.strip() for sh in lay.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip()
                     and not sh.is_placeholder]
            ph_text = next((p.text_frame.text.strip() for p in lay.placeholders
                            if p.has_text_frame and p.text_frame.text.strip()), "")
            special.append(
                f"[{i:2d}] {lay.name} — kept exactly as drawn. Its own wording "
                f"{ph_text or '(none)'!r}"
                + (f" plus {len(extra)} non-placeholder text box(es): "
                   f"{extra[0][:52]!r}" if extra else "")
                + '. Set "keep_closing": false on the slide to edit it.')
        elif low in ("agenda", "contents", "table of contents", "目錄"):
            bodies = [p for p in lay.placeholders
                      if "BODY" in str(p.placeholder_format.type)]
            special.append(
                f"[{i:2d}] {lay.name} — filled from the deck's own section headings "
                f"and their page numbers. It has {len(bodies)} body placeholder(s)"
                + ("; the rightmost takes the page numbers."
                   if len(bodies) >= 2 else "; headings and pages share one column."))
    if special:
        print("\n--- Layouts build_deck.py handles specially " + "-" * 18)
        for line in special:
            print(f"  {line}")

    long_titles = []
    for i, lay in enumerate(layouts):
        for p in lay.placeholders:
            if "TITLE" not in str(p.placeholder_format.type):
                continue
            szs = re.findall(r'sz="(\d+)"', p._element.xml)
            sz = int(szs[0]) / 100.0 if szs else 32.0
            w, h = inches(p.width), inches(p.height)
            # Chars that fit on one line, latin; CJK counts double.
            fit = int(8.6 * w * 0.94 * (18.0 / sz))
            long_titles.append((i, lay.name, sz, w, h, fit))
            break
    if long_titles:
        print("\n--- Headline budget (one line only) " + "-" * 26)
        print("  A headline must be one line: on the short boxes a second line overflows,")
        print("  and on the tall ones it fits but stops reading as a heading.")
        for i, name, sz, w, h, fit in long_titles:
            room = "1 line" if h < 2 * (sz * 1.45 / 72) else f"{h}in box"
            print(f"    [{i:2d}] {name:<30} {sz:g}pt in {w}in ({room})  "
                  f"~{fit} latin / ~{fit // 2} CJK chars")

    m = prs.slide_master
    print("\n--- Master placeholders (inches) " + "-" * 29)
    for p in m.placeholders:
        print(f"  {str(p.placeholder_format.type):<28} "
              f"x={inches(p.left)} y={inches(p.top)} "
              f"w={inches(p.width)} h={inches(p.height)}")

    print("\n--- Colours in slides/layouts/master " + "-" * 25)
    for hexv, n in colors.most_common(14):
        print(f"  #{hexv}  ×{n}")

    print("\n--- Fonts " + "-" * 52)
    for f, n in fonts.most_common(10):
        print(f"  {f:<24} ×{n}")
    print(f"  sizes present (pt): "
          f"{', '.join(str(s) for s in sorted(sizes, reverse=True))}")

    print("\n--- Theme (often NOT the real brand) " + "-" * 25)
    print(f"  latin fonts: {', '.join(sorted(set(theme_fonts))) or '—'}")
    print(f"  accent1:     {', '.join('#' + a.lower() for a in set(theme_accents)) or '—'}")
    if theme_fonts and {"Calibri", "Calibri Light"} & set(theme_fonts):
        print("  ^ stock Office theme. Read the brand from the colour/font counts above,")
        print("    not from here.")

    print("\nUse the layout NAMES above as the 'layout' value in the deck spec.\n")


if __name__ == "__main__":
    main()
